"""Tests for inrix_tools.reference (ROADMAP Item 29)."""
import datetime as dt
from pathlib import Path

import pandas as pd
import pytest

from inrix_tools import reference as ref

REPO_ROOT = Path(__file__).resolve().parents[1]
TT_LOGGER = REPO_ROOT / "TT Logger.xlsx"
TZ = "America/Denver"


# ---------------------------------------------------------------------------
# A synthetic workbook covering both header dialects
# ---------------------------------------------------------------------------
def _write_sheet(ws, origin, destination, days, start, end, samples):
    """One TT-Logger-shaped sheet: 5-row header block, blank row, column header,
    then ``(timestamp, seconds)`` samples."""
    for row, (label, value) in enumerate(
        [("Origin", origin), ("Destination", destination), ("Days", days),
         ("Start", start), ("End", end)], start=1
    ):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
    for col, name in enumerate(
        ["Timestamp", "Travel Time (s)", "Travel Time (min)", "Extra TT (Min)"], start=1
    ):
        ws.cell(row=7, column=col, value=name)
    for i, (stamp, secs) in enumerate(samples, start=8):
        ws.cell(row=i, column=1, value=stamp)
        ws.cell(row=i, column=2, value=secs)
        ws.cell(row=i, column=3, value=round(secs / 60.0, 6))
        ws.cell(row=i, column=4, value=0.5)


def _samples(start="2026-03-03 08:00:50", n=4, step_minutes=30, secs=300):
    t0 = pd.Timestamp(start)
    return [(t0 + pd.to_timedelta(step_minutes * i, unit="m"), secs + i) for i in range(n)]


@pytest.fixture
def workbook(tmp_path):
    """Two header dialects in one file: coordinates + ``datetime.time`` clocks on
    one sheet, a place name + an Excel **day fraction** (0.625 -> 15:00) on the
    other — both appear in the real workbook."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("Coord Route"),
                 "43.662633,-116.664553", "43.663139, -116.609317",
                 "Tue-Thur", dt.time(6, 0), dt.time(18, 0), _samples())
    _write_sheet(wb.create_sheet("Place Route"),
                 "Garden Valley, ID", "Horseshoe Bend, ID",
                 "Fri-Mon", 0.625, 0.7916666666666666,
                 _samples(start="2026-03-04 15:00:50", secs=1700))
    path = tmp_path / "toy_logger.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Header-block parsing
# ---------------------------------------------------------------------------
def test_parse_endpoint_distinguishes_coordinates_from_place_names():
    a = ref.parse_endpoint("43.662633,-116.664553")
    assert a["kind"] == "latlon" and a["lat"] == pytest.approx(43.662633)
    assert a["lon"] == pytest.approx(-116.664553)
    assert ref.parse_endpoint("43.66, -116.60")["kind"] == "latlon"   # space after comma
    # a place name also contains a comma — the parse is by value, not separator
    p = ref.parse_endpoint("Garden Valley, ID")
    assert p["kind"] == "place" and pd.isna(p["lat"])
    assert ref.parse_endpoint(None)["kind"] is None
    assert ref.parse_endpoint("   ")["kind"] is None


def test_parse_clock_handles_all_three_dialects():
    assert ref.parse_clock(dt.time(6, 30)) == dt.time(6, 30)
    assert ref.parse_clock(0.625) == dt.time(15, 0)          # Excel day fraction
    assert ref.parse_clock(0.7916666666666666) == dt.time(19, 0)
    assert ref.parse_clock("07:00") == dt.time(7, 0)
    assert ref.parse_clock("7:00 AM") == dt.time(7, 0)
    assert ref.parse_clock(pd.Timestamp("2026-03-03 18:45")) == dt.time(18, 45)
    assert ref.parse_clock(None) is None and ref.parse_clock(float("nan")) is None
    assert ref.parse_clock("not a time") is None


def test_route_table_marks_what_can_be_chain_matched(workbook):
    routes = ref.load_tt_logger_routes(workbook)
    coord = routes.set_index(ref.ROUTE_COL).loc["Coord Route"]
    place = routes.set_index(ref.ROUTE_COL).loc["Place Route"]

    assert coord["origin_kind"] == "latlon" and coord["destination_kind"] == "latlon"
    assert coord["origin_lat"] == pytest.approx(43.662633)
    assert coord["start_time"] == "06:00:00" and coord["end_time"] == "18:00:00"
    assert bool(coord[ref.CHAIN_MATCHABLE_COL])

    assert place["origin_kind"] == "place"
    assert place["start_time"] == "15:00:00" and place["end_time"] == "19:00:00"
    assert not bool(place[ref.CHAIN_MATCHABLE_COL])   # a geocoder's guess
    assert place["days"] == "Fri-Mon"


def test_route_endpoints_refuses_a_place_name_route(workbook):
    routes = ref.load_tt_logger_routes(workbook)
    (lat, lon), (lat2, lon2) = ref.route_endpoints(routes, "Coord Route")
    assert (lat, lon2) == pytest.approx((43.662633, -116.609317))
    with pytest.raises(ValueError, match="place names"):
        ref.route_endpoints(routes, "Place Route")
    with pytest.raises(KeyError):
        ref.route_endpoints(routes, "No Such Route")


def test_missing_timestamp_header_names_the_sheet(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Broken"
    wb.active["A1"] = "Origin"
    path = tmp_path / "broken.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="Broken"):
        ref.load_tt_logger_routes(path)


# ---------------------------------------------------------------------------
# Observations
# ---------------------------------------------------------------------------
def test_load_tt_logger_is_tidy_and_tz_aware(workbook):
    obs = ref.load_tt_logger(workbook, tz=TZ)
    assert list(obs.columns) == [ref.ROUTE_COL, "Date Time", ref.TT_COL, ref.EXTRA_TT_COL]
    assert str(obs["Date Time"].dt.tz) == TZ
    assert obs[ref.ROUTE_COL].nunique() == 2
    # minutes come from the seconds column (exact), not the rounded minutes column
    assert obs[ref.TT_COL].iloc[0] == pytest.approx(300 / 60.0)
    assert obs.attrs["units"] == {"travel_time": "Minutes"}
    assert obs.attrs["tz"] == TZ and obs.attrs["n_ambiguous_dropped"] == 0
    assert ref.routes_frame(obs).shape[0] == 2
    assert ref.load_tt_logger(workbook, sheets=["Place Route"])[ref.ROUTE_COL].unique().tolist() == ["Place Route"]


def test_dst_fold_samples_are_dropped_and_counted(tmp_path):
    """01:15 on 2025-11-02 happens twice in America/Denver. Guessing an offset
    would silently misplace an hour of data; the loader drops and counts it."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("Fold"), "43.0,-116.0", "43.1,-116.0", "Fri-Mon",
                 dt.time(1, 0), dt.time(3, 0),
                 [(pd.Timestamp("2025-11-02 01:15:49"), 300),
                  (pd.Timestamp("2025-11-02 03:15:49"), 310)])
    path = tmp_path / "fold.xlsx"
    wb.save(path)

    obs = ref.load_tt_logger(path, tz=TZ)
    assert len(obs) == 1 and obs.attrs["n_ambiguous_dropped"] == 1
    # the caller can override the policy explicitly
    kept = ref.load_tt_logger(path, tz=TZ, ambiguous=True)
    assert len(kept) == 2 and kept.attrs["n_ambiguous_dropped"] == 0


def test_tidy_frame_survives_a_parquet_round_trip(workbook, tmp_path):
    """``attrs`` rides along on ``to_parquet`` and is JSON-encoded, so the route
    table is stored as records — a DataFrame (or a ``datetime.time``) there makes
    the whole frame unwritable."""
    obs = ref.load_tt_logger(workbook, tz=TZ)
    path = tmp_path / "obs.parquet"
    obs.to_parquet(path)
    back = pd.read_parquet(path)
    assert len(back) == len(obs)
    assert ref.routes_frame(back)[ref.ROUTE_COL].tolist() == ["Coord Route", "Place Route"]


# ---------------------------------------------------------------------------
# Bin alignment: floor, not round
# ---------------------------------------------------------------------------
def test_align_to_bins_floors_where_round_would_disagree():
    """A sample at :58:50 belongs to the 15-minute bin that *contains* it (:45),
    not the one it is nearest to (:00 of the next hour). INRIX labels a bin by its
    start, so flooring is what the label means."""
    obs = pd.DataFrame({
        ref.ROUTE_COL: ["R", "R"],
        "Date Time": pd.to_datetime(["2026-03-03 08:58:50", "2026-03-03 09:03:20"]).tz_localize(TZ),
        ref.TT_COL: [5.0, 6.0],
    })
    binned = ref.align_to_bins(obs, bin_minutes=15)
    assert list(binned["Date Time"].dt.strftime("%H:%M")) == ["08:45", "09:00"]
    # rounding the *raw* samples would have put the first one in the 09:00 bin
    assert list(obs["Date Time"].dt.round("15min").dt.strftime("%H:%M")) == ["09:00", "09:00"]
    assert list(binned[ref.SAMPLE_TIME_COL].dt.strftime("%H:%M:%S")) == ["08:58:50", "09:03:20"]
    assert binned.attrs["bin_minutes"] == 15
    assert str(binned["Date Time"].dt.tz) == TZ


def test_align_to_bins_refuses_naive_timestamps():
    obs = pd.DataFrame({ref.ROUTE_COL: ["R"],
                        "Date Time": pd.to_datetime(["2026-03-03 08:58:50"]),
                        ref.TT_COL: [5.0]})
    with pytest.raises(ValueError, match="tz-aware"):
        ref.align_to_bins(obs)


# ---------------------------------------------------------------------------
# Gate 1 — nesting  (G3)
# ---------------------------------------------------------------------------
def _two_routes(sub_values, full_values, start="2026-03-03 08:00"):
    stamps = pd.date_range(start, periods=len(sub_values), freq="15min", tz=TZ)
    return pd.concat([
        pd.DataFrame({ref.ROUTE_COL: "Sub", "Date Time": stamps, ref.TT_COL: sub_values}),
        pd.DataFrame({ref.ROUTE_COL: "Full", "Date Time": stamps, ref.TT_COL: full_values}),
    ], ignore_index=True)


def test_nesting_gate_fires_on_an_impossible_sub_route():
    obs = _two_routes([6.4] * 10, [5.7] * 10)          # sub longer than its container
    got = ref.nesting_gate(obs, pairs=[("Sub", "Full")]).iloc[0]
    assert got["n_shared_bins"] == 10
    assert got["mean_diff"] == pytest.approx(0.7)
    assert got["frac_sub_exceeds_full"] == 1.0
    assert bool(got["violates"])


def test_nesting_gate_stays_quiet_on_a_clean_pair():
    obs = _two_routes([4.0] * 10, [5.5] * 10)
    got = ref.nesting_gate(obs, pairs=[("Sub", "Full")]).iloc[0]
    assert got["frac_sub_exceeds_full"] == 0.0 and not bool(got["violates"])
    assert got["mean_diff"] == pytest.approx(-1.5)

    # a couple of noisy bins are tolerated; a systematic reversal is not
    noisy = _two_routes([4.0] * 19 + [6.0], [5.5] * 20)
    assert not bool(ref.nesting_gate(noisy, pairs=[("Sub", "Full")]).iloc[0]["violates"])


def test_nested_pairs_from_chain_membership():
    """Nesting comes from the assembled chains, not from the sheet names — nothing
    in "Franklin WB - No Mid" says it is a subset of "Franklin WB"."""
    chains = {"Full": [1, 2, 3, 4], "Sub": [2, 3], "Other": [7, 8], "Same": [1, 2, 3, 4]}
    assert ref.nested_pairs(chains) == [("Sub", "Full"), ("Sub", "Same")]

    obs = _two_routes([6.4] * 6, [5.7] * 6)
    got = ref.nesting_gate(obs, chains={"Sub": [2, 3], "Full": [1, 2, 3, 4]})
    assert got.loc[0, "sub_route"] == "Sub" and bool(got.loc[0, "violates"])
    with pytest.raises(ValueError, match="pairs"):
        ref.nesting_gate(obs)


def test_nesting_gate_compares_only_shared_bins():
    obs = _two_routes([6.4] * 10, [5.7] * 10)
    obs = obs.drop(obs[(obs[ref.ROUTE_COL] == "Full")].index[:4])
    assert ref.nesting_gate(obs, pairs=[("Sub", "Full")]).iloc[0]["n_shared_bins"] == 6


# ---------------------------------------------------------------------------
# Gate 2 — coverage  (G4)
# ---------------------------------------------------------------------------
def test_coverage_gate_exposes_a_short_route():
    """The G4 shape: one route logs all year, another stops in January. Both sit
    under the same study-window banner until this frame separates them."""
    long_stamps = pd.date_range("2026-01-02 08:00", "2026-08-20 08:00", freq="D", tz=TZ)
    short_stamps = pd.date_range("2026-01-02 08:00", "2026-01-17 08:00", freq="D", tz=TZ)
    obs = pd.concat([
        pd.DataFrame({ref.ROUTE_COL: "Long", "Date Time": long_stamps, ref.TT_COL: 5.0}),
        pd.DataFrame({ref.ROUTE_COL: "Short", "Date Time": short_stamps, ref.TT_COL: 5.0}),
    ], ignore_index=True)

    got = ref.coverage_gate(obs, window=("2026-01-01", "2026-08-24")).set_index(ref.ROUTE_COL)
    assert got.loc["Short", "n_days"] == 16
    assert got.loc["Short", "days_in_window"] == 16
    assert not bool(got.loc["Short", "covers_window"])
    assert got.loc["Short", "window_covered_fraction"] < 0.07
    assert bool(got.loc["Long", "covers_window"]) is False   # stops 4 days early
    assert got.loc["Long", "n_days"] == 231
    assert got.loc["Long", "sampled_day_fraction"] == pytest.approx(1.0)


def test_coverage_gate_without_a_window():
    stamps = pd.to_datetime(["2026-01-02 08:00", "2026-01-02 09:00", "2026-01-06 08:00"]).tz_localize(TZ)
    obs = pd.DataFrame({ref.ROUTE_COL: "R", "Date Time": stamps, ref.TT_COL: 5.0})
    got = ref.coverage_gate(obs).iloc[0]
    assert got["n_obs"] == 3 and got["n_days"] == 2 and got["span_days"] == 5
    assert got["sampled_day_fraction"] == pytest.approx(0.4)
    assert "covers_window" not in got.index


def test_coverage_gate_measures_the_sampling_cadence():
    """The report must state the cadence it measured, not the one it assumed: the
    outside pass published "15-minute intervals" for a reference logged about
    twice an hour (G10). Two samples an hour over three days, and an overnight
    gap that must not move the median."""
    stamps = []
    for day in ("2026-01-02", "2026-01-03", "2026-01-04"):
        stamps += list(pd.date_range(f"{day} 06:00", f"{day} 10:00", freq="30min", tz=TZ))
    obs = pd.DataFrame({ref.ROUTE_COL: "R", "Date Time": stamps, ref.TT_COL: 5.0})

    got = ref.coverage_gate(obs).iloc[0]
    assert got["median_sample_minutes"] == pytest.approx(30.0)
    assert got["obs_per_day"] == pytest.approx(9.0)


# ---------------------------------------------------------------------------
# Gate 3 — clock agreement  (G9)
# ---------------------------------------------------------------------------
def _diurnal(start="2026-03-03 06:00", periods=96, amp=3.0, base=10.0, route="R"):
    stamps = pd.date_range(start, periods=periods, freq="15min", tz=TZ)
    hours = stamps.hour + stamps.minute / 60.0
    import numpy as np
    return pd.DataFrame({ref.ROUTE_COL: route, "Date Time": stamps,
                         ref.TT_COL: base + amp * np.sin(2 * np.pi * hours / 24.0)})


def test_lag_scan_finds_zero_when_the_clocks_agree():
    reference = _diurnal()
    got = ref.lag_summary(ref.lag_scan(reference, reference.copy())).iloc[0]
    assert got["best_lag_minutes"] == 0 and bool(got["agrees_at_zero"])
    assert got["sd_at_zero"] == pytest.approx(0.0, abs=1e-9)


def test_lag_scan_finds_an_injected_offset():
    reference = _diurnal()
    shifted = reference.copy()
    shifted["Date Time"] = shifted["Date Time"] + pd.to_timedelta(30, unit="m")
    got = ref.lag_summary(ref.lag_scan(reference, shifted)).iloc[0]
    assert got["best_lag_minutes"] == 30
    assert not bool(got["agrees_at_zero"])
    assert got["improvement"] > 0.5


def test_lag_scan_statistic_ignores_a_constant_bias():
    """A lag scan reads shape, not level. Scanning on raw RMSE lets a large
    systematic bias (which every lag carries) swamp the diurnal signal — on the
    real corridors that moved the apparent best lag to +30/+60 on half the routes.
    ``sd_diff`` is invariant to the offset; ``rmse`` is not."""
    reference = _diurnal()
    biased = reference.copy()
    biased[ref.TT_COL] = biased[ref.TT_COL] + 8.0

    plain = ref.lag_scan(reference, reference.copy())
    with_bias = ref.lag_scan(reference, biased)
    merged = plain.merge(with_bias, on=[ref.ROUTE_COL, "lag_minutes"], suffixes=("", "_b"))
    assert merged["sd_diff"].to_numpy() == pytest.approx(merged["sd_diff_b"].to_numpy())
    assert (merged["rmse_b"] > merged["rmse"] + 1.0).all()
    assert ref.lag_summary(with_bias).iloc[0]["best_lag_minutes"] == 0


def test_lag_scan_reports_every_lag_and_marks_the_best():
    scan = ref.lag_scan(_diurnal(), _diurnal(), max_lag_minutes=30, step_minutes=15)
    assert sorted(scan["lag_minutes"].unique()) == [-30, -15, 0, 15, 30]
    assert scan.loc[scan["is_best"], "lag_minutes"].tolist() == [0]
    assert (scan["n"] > 0).all()


# ---------------------------------------------------------------------------
# The real workbook (skipped when it isn't present)
# ---------------------------------------------------------------------------
@pytest.mark.skipif(not TT_LOGGER.exists(), reason="TT Logger workbook not available")
def test_real_workbook_reproduces_the_franklin_nesting_violation():
    """G3, measured: Google's "Franklin WB - No Mid" exceeds its own full
    "Franklin WB" in >99% of shared bins — while the EB pair is clean, so this is
    the reference's error, not a systematic artifact of the check."""
    sheets = ["Franklin EB", "Franklin WB", "Franklin EB - No Mid", "Franklin WB - No Mid"]
    obs = ref.align_to_bins(ref.load_tt_logger(TT_LOGGER, sheets=sheets))
    got = ref.nesting_gate(obs, pairs=[("Franklin WB - No Mid", "Franklin WB"),
                                       ("Franklin EB - No Mid", "Franklin EB")]).set_index("sub_route")

    wb = got.loc["Franklin WB - No Mid"]
    assert wb["n_shared_bins"] > 4000
    assert wb["frac_sub_exceeds_full"] > 0.99 and bool(wb["violates"])
    assert wb["mean_sub"] > wb["mean_full"]
    eb = got.loc["Franklin EB - No Mid"]
    assert eb["frac_sub_exceeds_full"] == 0.0 and not bool(eb["violates"])


@pytest.mark.skipif(not TT_LOGGER.exists(), reason="TT Logger workbook not available")
def test_real_workbook_header_dialects_and_dst_fold():
    routes = ref.load_tt_logger_routes(TT_LOGGER)
    assert set(routes["origin_kind"]) == {"latlon", "place"}
    assert (~routes[ref.CHAIN_MATCHABLE_COL]).sum() == 3      # the three city-to-city sheets
    # "VSL NB PM" stores its window as an Excel day fraction, not a time
    assert routes.set_index(ref.ROUTE_COL).loc["VSL NB PM", "start_time"] == "15:00:00"

    obs = ref.load_tt_logger(TT_LOGGER, sheets=["Franklin EB"])
    assert obs.attrs["n_ambiguous_dropped"] == 4              # the 2025-11-02 fold
